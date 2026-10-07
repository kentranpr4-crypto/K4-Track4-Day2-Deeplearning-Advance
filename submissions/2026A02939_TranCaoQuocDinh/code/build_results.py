"""Build submission tables from recorded runs and predictions; never invent metrics."""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from eval import CLASS_NAMES, compute_metrics, mean_std, read_pred


def _prediction_metrics(path):
    prediction = read_pred(str(path))
    return compute_metrics(prediction.y_true, prediction.y_pred, prediction.probs)


def _sort_id(path):
    match = re.search(r"seed(\d+)", str(path))
    return int(match.group(1)) if match else -1


def make_test_figures(root: Path, images_dir: Path | None = None):
    prediction_path = root / "predictions/F01_seed0_test.csv"
    if not prediction_path.exists():
        return
    import matplotlib.pyplot as plt
    import numpy as np
    from PIL import Image
    from eval import confusion_matrix

    prediction = read_pred(str(prediction_path))
    figures = root / "figures"
    figures.mkdir(exist_ok=True)
    cm = confusion_matrix(prediction.y_true, prediction.y_pred)
    fig, ax = plt.subplots(figsize=(9, 8))
    image = ax.imshow(cm, cmap="Blues")
    ax.set_xticks(range(9), CLASS_NAMES, rotation=45, ha="right")
    ax.set_yticks(range(9), CLASS_NAMES)
    ax.set(xlabel="Predicted", ylabel="True", title="F01 seed 0 - test confusion matrix")
    for row in range(9):
        for col in range(9):
            ax.text(col, row, str(cm[row, col]), ha="center", va="center", fontsize=8)
    fig.colorbar(image, ax=ax)
    fig.tight_layout()
    fig.savefig(figures / "confusion_matrix_F01.png", dpi=160)
    plt.close(fig)

    run_root = root / "runs" if (root / "runs").is_dir() else root / "logs"
    config_path = run_root / "F01/seed0/config.json"
    if not config_path.exists():
        return
    images_dir = images_dir or Path(json.loads(config_path.read_text())["images_dir"])
    if not images_dir.is_absolute():
        images_dir = root / images_dir
    if not images_dir.is_dir():
        print(f"Error-image figure skipped: image directory missing ({images_dir})")
        return
    errors = np.flatnonzero(prediction.y_true != prediction.y_pred)
    # Show the two difficult classes first, then other errors.
    order = sorted(errors, key=lambda index: (prediction.y_true[index] not in (0, 7), index))[:9]
    fig, axes = plt.subplots(3, 3, figsize=(10, 10))
    for ax, index in zip(axes.flat, order):
        filename = prediction.filenames[index]
        with Image.open(images_dir / filename) as image:
            ax.imshow(image.convert("RGB"))
        ax.set_title(f"{CLASS_NAMES[prediction.y_true[index]]} -> {CLASS_NAMES[prediction.y_pred[index]]}", fontsize=8)
        ax.axis("off")
    for ax in list(axes.flat)[len(order):]:
        ax.axis("off")
    fig.tight_layout()
    fig.savefig(figures / "error_examples_F01.png", dpi=140)
    plt.close(fig)


def make_tradeoff_figure(root: Path, rows: list[dict]):
    if not rows:
        return
    import matplotlib.pyplot as plt

    figures = root / "figures"
    figures.mkdir(exist_ok=True)
    fig, ax = plt.subplots(figsize=(8, 5))
    for row in rows:
        if row.get("p95_ms") is None:
            continue
        ax.scatter(row["p95_ms"], row["macro_f1_val"], s=55)
        ax.annotate(row["method"].split("_")[0], (row["p95_ms"], row["macro_f1_val"]),
                    xytext=(5, 4), textcoords="offset points")
    ax.set(xlabel="p95 forward latency, batch 1 (ms)", ylabel="Validation macro-F1",
           title="Inference quality and latency on Tesla T4")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(figures / "inference_tradeoff.png", dpi=160)
    plt.close(fig)

def make_backbone_figure(root: Path, rows: list[dict]):
    if not rows:
        return
    import matplotlib.pyplot as plt

    figures = root / "figures"
    figures.mkdir(exist_ok=True)
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for row in rows:
        for ax, field in zip(axes, ("latency_batch1_ms", "params_m")):
            if row.get(field) is not None:
                ax.scatter(row[field], row["macro_f1_val"], s=55)
                ax.annotate(row["exp_id"], (row[field], row["macro_f1_val"]),
                            xytext=(4, 4), textcoords="offset points")
    axes[0].set(xlabel="p50 FP32 forward latency, batch 1 (ms)", ylabel="Validation macro-F1")
    axes[1].set(xlabel="Parameters (millions)", ylabel="Validation macro-F1")
    for ax in axes:
        ax.grid(alpha=0.25)
    fig.suptitle("Backbone comparison, fold 0 / seed 0")
    fig.tight_layout()
    fig.savefig(figures / "backbone_tradeoff.png", dpi=160)
    plt.close(fig)


def recipe_description(cfg):
    return (f"{cfg['backbone']}; init={cfg['init']}; aug={cfg['aug']}; "
            f"loss={cfg['loss']}; mix={cfg.get('mix')}; epochs={cfg['epochs']}; "
            f"batch={cfg['batch_size']}; lr={cfg['lr_backbone']}/{cfg['lr_head']}")


def collect(root: Path):
    backbones, training, final, per_class, latency, inference, summary = [], [], [], [], [], [], []
    runs = root / "runs" if (root / "runs").is_dir() else root / "logs"
    predictions = root / "predictions"
    baseline_path = runs / "B03/seed0/config.json"
    baseline_cfg = json.loads(baseline_path.read_text()) if baseline_path.exists() else None
    axes = {"T01": "A initialization", "T02": "B augmentation", "T03": "B augmentation",
            "T04": "C loss", "T05": "C loss", "T06": "C loss", "T07": "B+C combination"}
    for config_file in sorted(runs.glob("*/seed*/config.json")):
        run_dir = config_file.parent
        cfg = json.loads(config_file.read_text())
        exp_id, seed = cfg["exp_id"], cfg["seed"]
        history_path = run_dir / "history.csv"
        history = pd.read_csv(history_path) if history_path.exists() else pd.DataFrame()
        if not history.empty:
            curve_path = root / "curves" / f"{exp_id}_seed{seed}_{cfg['backbone']}.png"
            if not curve_path.exists():
                try:
                    from .train import plot_curves
                except ImportError:
                    from train import plot_curves
                plot_curves(history.to_dict("records"), curve_path,
                            f"{exp_id} seed {seed} - {cfg['backbone']}")
        val_path = predictions / f"{exp_id}_seed{seed}_val.csv"
        if not val_path.exists():
            continue
        val = _prediction_metrics(val_path)
        info_path = run_dir / "summary.json"
        info = json.loads(info_path.read_text()) if info_path.exists() else {}
        profile_path = run_dir / "profile.json"
        profile = json.loads(profile_path.read_text()) if profile_path.exists() else {}
        if profile.get("latency"):
            latency.append({"exp_id": exp_id, "seed": seed,
                            "method": "I00_one_view", "bn_fused": False,
                            "preprocessing_included": False, **profile["latency"]})
        tag_path = run_dir / "pretrained_tag.json"
        tag = json.loads(tag_path.read_text()).get("hf_hub_id", "") if tag_path.exists() else ""
        common = {"exp_id": exp_id, "backbone": cfg["backbone"], "seed": seed,
                  "macro_f1_val": val["macro_f1"], "top1_val": val["top1"],
                  "best_epoch": info.get("best_epoch", int(history.loc[history.val_macro_f1.idxmax(), "epoch"]) if not history.empty else None),
                  "params_m": info.get("params_m", profile.get("params_m")),
                  "gmac": info.get("gmac") or profile.get("gmac"),
                  "latency_batch1_ms": (profile.get("latency") or {}).get("p50"),
                  "train_seconds_per_epoch": history.epoch_seconds.mean() if "epoch_seconds" in history else None}
        if seed == 0 and not exp_id.startswith("F"):
            summary.append({**common, "ece_val": val["ece"], "configuration": recipe_description(cfg),
                            "method": "I00_one_view"})
        if exp_id.startswith("B"):
            backbones.append({**common, "pretrained_tag": tag, "img_size": cfg["img_size"],
                              "epochs": cfg["epochs"], "notes": "1 seed; pretrained tags differ"})
        if exp_id.startswith("T"):
            compare = baseline_cfg if exp_id != "T00" else None
            varied = ("init", "aug", "loss", "label_smoothing", "focal_gamma", "class_weight_beta")
            changes = "; ".join(f"{key}: {compare.get(key)} -> {cfg.get(key)}"
                                for key in varied if compare and compare.get(key) != cfg.get(key))
            training.append({**common, "init": cfg["init"], "aug": cfg["aug"],
                             "loss": cfg["loss"], "mix": cfg.get("mix"),
                             "sampler": cfg.get("sampler"), "ema_decay": cfg.get("ema_decay"),
                             "axis": axes.get(exp_id, "final reference"),
                             "baseline_exp_id": "B03" if compare else "T00",
                             "changes_from_baseline": changes or "reference recipe",
                             "chinee_apple_f1_val": float(val["f1"][0]),
                             "snake_weed_f1_val": float(val["f1"][7]),
                             "notes": "B03 uses baseline recipe on selected backbone; T00 is ResNet50 final reference"})
        test_path = predictions / f"{exp_id}_seed{seed}_test.csv"
        if test_path.exists():
            test = _prediction_metrics(test_path)
            inference_path = run_dir / "final_inference.json"
            method = json.loads(inference_path.read_text()).get("method") if inference_path.exists() else None
            final.append({**common, "macro_f1_test": test["macro_f1"],
                          "top1_test": test["top1"], "ece_test": test["ece"],
                          "inference_method": method, "configuration": recipe_description(cfg),
                          "val_inference_method": "I00_one_view",
                          "notes": "F01 val is 1-view; F01 test is 5-crop" if exp_id == "F01" else "val/test: 1-view"})
            for index, name in enumerate(CLASS_NAMES):
                per_class.append({"exp_id": exp_id, "seed": seed, "class": name,
                    "support": int(test["support"][index]), "precision": test["precision"][index],
                    "recall": test["recall"][index], "f1": test["f1"][index]})
    for path in sorted((root / "inference_out").glob("*_val_inference.json")):
        data = json.loads(path.read_text())
        baseline_p50 = (data["methods"].get("I00_one_view", {}).get("latency") or {}).get("p50")
        for method, record in data["methods"].items():
            metrics = record["metrics"]
            timing = record.get("latency") or {}
            inference.append({"exp_id": data["exp_id"], "seed": data["seed"],
                "method": method, "macro_f1_val": metrics["macro_f1"],
                "top1_val": metrics["top1"], "ece_val": metrics["ece"],
                "k": timing.get("k_views", 1), "p50_ms": timing.get("p50"),
                "p95_ms": timing.get("p95"), "p99_ms": timing.get("p99"),
                "images_per_s": timing.get("images_per_s"),
                "relative_p50": timing.get("p50") / baseline_p50 if timing.get("p50") and baseline_p50 else None,
                "gpu": timing.get("gpu"), "dtype": timing.get("dtype"),
                "batch": timing.get("batch"), "preprocessing_included": False})
            inference[-1]["checkpoint"] = f"{runs.name}/{data['exp_id']}/seed{data['seed']}/best.pt (external weights)"
            inference[-1]["checkpoint_epoch"] = next((row["best_epoch"] for row in backbones
                if row["exp_id"] == data["exp_id"] and row["seed"] == data["seed"]), None)
            summary.append({"exp_id": method.split("_")[0], "source_exp_id": data["exp_id"],
                "seed": data["seed"], "method": method, "macro_f1_val": metrics["macro_f1"],
                "top1_val": metrics["top1"], "ece_val": metrics["ece"],
                "latency_batch1_ms": timing.get("p50"), "p95_ms": timing.get("p95"),
                "relative_p50": inference[-1]["relative_p50"],
                "configuration": f"{data['exp_id']} checkpoint + {method}"})
            if timing:
                latency.append({"exp_id": data["exp_id"], "seed": data["seed"],
                                "method": method, "bn_fused": False,
                                "preprocessing_included": False,
                                "measurement_scope": "repeated forward passes; crop generation / aggregation excluded",
                                **timing})
    baseline = next((row["macro_f1_val"] for row in summary
                     if row["exp_id"] == "B03" and row["seed"] == 0), None)
    for row in training:
        row["delta_macro_f1_vs_B03"] = (row["macro_f1_val"] - baseline
                                        if baseline is not None and row["exp_id"] != "T00" else None)
    if final:
        final_frame = pd.DataFrame(final)
        for exp_id, group in final_frame.groupby("exp_id"):
            if len(group) < 2:
                continue
            source_row = next(row for row in final if row["exp_id"] == exp_id)
            aggregate = {"exp_id": exp_id, "seed": "mean +/- std", "seed_count": len(group),
                         "inference_method": source_row["inference_method"],
                         "configuration": source_row["configuration"],
                         "val_inference_method": "I00_one_view"}
            for field in ("macro_f1_val", "macro_f1_test", "top1_test", "ece_test"):
                avg, std = mean_std(group[field].tolist())
                aggregate[field] = f"{avg:.4f} +/- {std:.4f}"
            final.append(aggregate)
        class_frame = pd.DataFrame(per_class)
        for (exp_id, class_name), group in class_frame.groupby(["exp_id", "class"]):
            if len(group) < 2:
                continue
            aggregate = {"exp_id": exp_id, "seed": "mean +/- std", "class": class_name,
                         "support": int(group["support"].iloc[0])}
            for field in ("precision", "recall", "f1"):
                avg, std = mean_std(group[field].tolist())
                aggregate[field] = f"{avg:.4f} +/- {std:.4f}"
            per_class.append(aggregate)
    ranked_summary = sorted(summary, key=lambda row: row["macro_f1_val"], reverse=True)[:10]
    for rank, row in enumerate(ranked_summary, 1):
        row["rank"] = rank
    baseline_reference = next((row for row in summary if row["exp_id"] == "B01"), None)
    if baseline_reference and baseline_reference not in ranked_summary:
        ranked_summary.append({**baseline_reference, "rank": "reference baseline"})
    return {"Backbones": backbones, "Training": training, "Inference": inference,
            "Final": final, "PerClass": per_class, "Latency": latency,
            "Summary": ranked_summary}


def write_report(path: Path, sheets: dict):
    def markdown_table(frame):
        columns = list(frame.columns)
        def cell(value):
            return str(value).replace("|", "\\|").replace("\n", " ")
        rows = ["| " + " | ".join(map(cell, columns)) + " |",
                "| " + " | ".join("---" for _ in columns) + " |"]
        rows.extend("| " + " | ".join(cell(value) for value in row) + " |"
                    for row in frame.itertuples(index=False, name=None))
        return "\n".join(rows)

    lines = ["# DeepWeeds Lab Day 2 - Bao cao ket qua", "",
             "## Trang thai", "", "Bao cao duoc tao tu log va prediction that; cac muc chua co so lieu can duoc hoan thien sau khi chay Kaggle.", ""]
    backbones = pd.DataFrame(sheets["Backbones"])
    if not backbones.empty:
        lines += ["## Backbone tren validation", "", markdown_table(backbones), ""]
    for title, key in (("Ablation huan luyen", "Training"),
                       ("Phuong phap suy luan", "Inference"),
                       ("Ket qua test", "Final"),
                       ("Chi so tung lop", "PerClass")):
        frame = pd.DataFrame(sheets[key])
        lines += [f"## {title}", "", markdown_table(frame) if not frame.empty else "Chua co ket qua.", ""]
    final = pd.DataFrame([row for row in sheets["Final"] if isinstance(row.get("seed"), int)])
    if not final.empty:
        lines += ["## Mean va std theo seed", ""]
        for exp_id, group in final.groupby("exp_id"):
            lines.append(f"### {exp_id} ({len(group)} seed)")
            for field in ("macro_f1_test", "top1_test", "ece_test"):
                avg, std = mean_std(group[field].tolist())
                lines.append(f"- {field}: {avg:.4f} +/- {std:.4f}")
            lines.append("")
    if {"F01", "T00"} <= set(final.get("exp_id", [])):
        means = final.groupby("exp_id")["macro_f1_test"].mean()
        lines += ["## Cai thien so voi moc", "",
                  f"Delta macro-F1 test F01 - T00: {means['F01'] - means['T00']:+.4f}.", ""]
    if (path.parent / "figures" / "confusion_matrix_F01.png").exists():
        lines += ["## Phan tich loi", "", "![Confusion matrix](figures/confusion_matrix_F01.png)", "",
                  "![Error examples](figures/error_examples_F01.png)", ""]
    lines += ["## Phan tich va han che can bo sung", "",
              "- Giai thich moi thay doi dua tren ket qua validation va do nhieu qua seed.",
              "- Them ma tran nham lan, anh du doan sai va khuyen nghi latency cho robot.",
              "- Neu gioi han: mot fold, chia ngau nhien theo anh thay vi dia diem, co the lac quan tren dia diem moi.", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path("."))
    parser.add_argument("--xlsx", type=Path, default=Path("results.xlsx"))
    parser.add_argument("--report", type=Path, default=Path("report_draft.md"))
    parser.add_argument("--images-dir", type=Path, help="Image directory for error examples")
    parser.add_argument("--skip-draft", action="store_true", help="Preserve a manually completed report")
    args = parser.parse_args()
    sheets = collect(args.root)
    make_test_figures(args.root, args.images_dir.resolve() if args.images_dir else None)
    make_tradeoff_figure(args.root, sheets["Inference"])
    make_backbone_figure(args.root, sheets["Backbones"])
    from openpyxl.styles import Alignment, Font, PatternFill
    with pd.ExcelWriter(args.xlsx, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            frame = pd.DataFrame(rows)
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = "1:1"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="243B53")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 45
            numeric_f1 = [row for row in rows if isinstance(row.get("macro_f1_val"), (float, int))]
            best_f1 = max((row["macro_f1_val"] for row in numeric_f1), default=None)
            for index, row in enumerate(rows, start=2):
                for cell in ws[index]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"
                    if (best_f1 is not None and row.get("macro_f1_val") == best_f1) or row.get("seed") == "mean +/- std":
                        cell.fill = PatternFill("solid", fgColor="E3F2DF")
                ws.row_dimensions[index].height = 42
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(38, max(12,
                    max(len(str(cell.value or "")) for cell in column) + 2))
    if not args.skip_draft:
        write_report(args.report, sheets)
    print(f"Saved {args.xlsx}" + (f" and {args.report}" if not args.skip_draft else ""))


if __name__ == "__main__":
    main()
